import subprocess
import sys
import tempfile
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

RAIZ = Path(__file__).resolve().parent.parent
SCRIPT = RAIZ / "scripts" / "gastos.py"
EJEMPLO = RAIZ / "assets" / "ejemplo.csv"
EXCEL = RAIZ / "assets" / "ejemplo-excel.csv"
CODIGO_OK = 0
CODIGO_DEPENDENCIA = 3
PASA = "PASA"
FALLA = "FALLA"
OMITIDO = "OMITIDO"

CASOS = []


def registrar(nombre):
    def envoltura(funcion):
        CASOS.append((nombre, funcion))
        return funcion

    return envoltura


def correr(argumentos):
    return subprocess.run(
        [sys.executable, str(SCRIPT), *argumentos],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )


def escribir(directorio, nombre, contenido):
    ruta = directorio / nombre
    ruta.write_text(contenido, encoding="utf-8")
    return ruta


def falla(proceso, esperado, motivo):
    return f"exit {proceso.returncode} != {esperado} o la salida no menciona {motivo!r}"


@registrar("Caso 1 - Exito: clasifica y resume los 15 gastos de ejemplo")
def caso_exito(directorio):
    salida = directorio / "resumen.csv"
    proceso = correr([str(EJEMPLO), "-o", str(salida)])
    if proceso.returncode != CODIGO_OK:
        return FALLA, falla(proceso, CODIGO_OK, "POR CATEGORIA")
    if not salida.exists():
        return FALLA, f"no se genero {salida.name}"
    if "Total   : Bs 3,162.80" not in proceso.stdout:
        return FALLA, "el total no coincide con la suma esperada (Bs 3,162.80)"
    if "Regalo para mi mama" not in proceso.stdout:
        return FALLA, "no reporto el gasto sin regla explicita"
    return PASA, "15 gastos clasificados, total correcto y 1 marcado para revisar"


@registrar("Caso 2 - Entrada invalida: falta la columna 'monto'")
def caso_cabecera(directorio):
    ruta = escribir(
        directorio,
        "sin_monto.csv",
        "fecha,descripcion\n2026-09-01,Supermercado Maxi\n",
    )
    proceso = correr([str(ruta)])
    if proceso.returncode != 1:
        return FALLA, falla(proceso, 1, "Faltan columnas requeridas")
    if "Faltan columnas requeridas: monto" not in proceso.stderr:
        return FALLA, "no dijo que columna falta"
    return PASA, "rechaza el archivo y lista las columnas requeridas"


@registrar("Caso 3 - Entrada invalida: monto que no es numero")
def caso_monto(directorio):
    ruta = escribir(
        directorio,
        "monto_malo.csv",
        "fecha,descripcion,monto\n2026-09-01,Almuerzo,doce\n2026-09-02,Taxi,18\n",
    )
    proceso = correr([str(ruta)])
    if proceso.returncode != 1:
        return FALLA, falla(proceso, 1, "monto invalido")
    if "Fila 2: monto invalido 'doce'" not in proceso.stderr:
        return FALLA, "no indico la fila ni el valor culpable"
    return PASA, "detecta la fila 2, la senala y no genera resumen"


@registrar("Caso 4 - Entrada invalida: fecha y categoria desconocidas")
def caso_fecha_y_categoria(directorio):
    ruta = escribir(
        directorio,
        "basura.csv",
        "fecha,descripcion,monto,categoria\n"
        "ayer,Panaderia,35,comida\n"
        "2026-09-02,Taxi,18,transporte\n",
    )
    proceso = correr([str(ruta)])
    if proceso.returncode != 1:
        return FALLA, falla(proceso, 1, "fecha invalida")
    errores = proceso.stderr
    if "fecha invalida 'ayer'" not in errores or "categoria invalida 'transporte'" not in errores:
        return FALLA, "no reporto los dos errores"
    return PASA, "acumula los 2 errores de una vez en vez de fallar en el primero"


@registrar("Caso 5 - Problema habitual: archivo que no existe")
def caso_inexistente(directorio):
    proceso = correr([str(directorio / "no_existe.csv")])
    if proceso.returncode != 2:
        return FALLA, falla(proceso, 2, "No se encontro el archivo")
    return PASA, "distingue 'archivo faltante' (codigo 2) de 'contenido invalido' (codigo 1)"


@registrar("Caso 6 - Entrada real de Excel en Bolivia: separador ';' y montos 1.240,80")
def caso_excel(directorio):
    salida = directorio / "resumen-excel.csv"
    proceso = correr([str(EXCEL), "-o", str(salida)])
    if proceso.returncode != CODIGO_OK:
        return FALLA, falla(proceso, CODIGO_OK, "POR CATEGORIA")
    if "Total   : Bs 3,162.80" not in proceso.stdout:
        return FALLA, "el total no coincide con el del CSV estandar (Bs 3,162.80)"
    if "Regalo para mi mama" not in proceso.stdout:
        return FALLA, "no reporto el gasto sin regla explicita"
    return PASA, "mismos 15 gastos y mismo total, leyendo ';' y cabeceras con tilde"


@registrar("Caso 7 - Salida XLSX organizada con hojas y detalle")
def caso_xlsx(directorio):
    salida = directorio / "resumen.xlsx"
    proceso = correr([str(EXCEL), "-o", str(salida), "-d"])
    if proceso.returncode == CODIGO_DEPENDENCIA:
        return OMITIDO, "falta openpyxl; instala con assets/requirements-xlsx.txt"
    if proceso.returncode != CODIGO_OK:
        return FALLA, falla(proceso, CODIGO_OK, "Resumen escrito en")
    if not salida.exists():
        return FALLA, f"no se genero {salida.name}"

    try:
        from openpyxl import load_workbook
    except ImportError:
        return OMITIDO, "falta openpyxl; instala con assets/requirements-xlsx.txt"

    libro = load_workbook(salida, data_only=True)
    esperadas = {"Resumen", "Por subcategoria", "Por categoria", "Revisar", "Detalle"}
    if set(libro.sheetnames) != esperadas:
        return FALLA, f"hojas inesperadas: {libro.sheetnames}"
    if libro["Resumen"]["B5"].value != 3162.8:
        return FALLA, "el total del resumen XLSX no coincide"
    if libro["Revisar"].max_row != 4:
        return FALLA, "la hoja Revisar no contiene exactamente 1 gasto"
    if libro["Detalle"].max_row != 18:
        return FALLA, "la hoja Detalle no contiene los 15 gastos"
    return PASA, "5 hojas formateadas, total correcto y 1 gasto para revisar"


def main():
    print()
    print("=" * 78)
    print("DEMOSTRACION DE control-gastos")
    print("=" * 78)
    print(f"Python  : {sys.version.split()[0]}")
    print(f"Script  : {SCRIPT.relative_to(RAIZ.parent)}")
    print(f"Entradas: {EJEMPLO.relative_to(RAIZ.parent)} y {EXCEL.relative_to(RAIZ.parent)}")
    print()

    resultados = []
    with tempfile.TemporaryDirectory() as temporal:
        directorio = Path(temporal)
        for numero, (nombre, funcion) in enumerate(CASOS, start=1):
            print("-" * 78)
            print(f"[{numero}] {nombre}")
            try:
                estado, detalle = funcion(directorio)
            except Exception as error:
                estado, detalle = FALLA, f"excepcion inesperada: {error}"
            resultados.append(estado)
            print(f"    {estado:<8} - {detalle}")

    print("-" * 78)
    pasado = resultados.count(PASA)
    omitido = resultados.count(OMITIDO)
    fallado = resultados.count(FALLA)
    if fallado:
        linea = f"RESULTADO: {pasado}/{len(resultados)} casos pasan - hay fallas que corregir"
    elif omitido:
        linea = f"RESULTADO: {pasado}/{len(resultados)} casos pasan, {omitido} omitido (falta openpyxl)"
    else:
        linea = f"RESULTADO: {pasado}/{len(resultados)} casos pasan"
    print(linea)
    print("=" * 78)
    print()
    return 1 if fallado else 0


if __name__ == "__main__":
    sys.exit(main())
